"""
Builds the synthetic source-model workbooks used by tests/source-workbook.test.ts.

These are NOT the Avignon files. They mirror the patterns described for the
real model (defined names, a table, formulas with cached results, percent and
currency formats, a sheet-scoped duplicate name, a range name, a hidden sheet)
at a size a deterministic test can use.

Run from the repo root:
    python3 tests/fixtures/build-source-model-fixture.py
Then the .xlsx is recalculated with LibreOffice so every formula has a cached
value, exactly as Excel would store it.
"""
import os
import shutil
import subprocess
import zipfile

import openpyxl
from openpyxl.styles import numbers
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "source-model.xlsx")
XLSM = os.path.join(HERE, "source-model.xlsm")
RECALC = "/mnt/skills/public/xlsx/scripts/recalc.py"

CURRENCY = '"$"#,##0_);("$"#,##0)'
PERCENT = "0.00%"
PERCENT1 = "0.0%"

wb = openpyxl.Workbook()

# ---- Assumptions -------------------------------------------------------
assumptions = wb.active
assumptions.title = "Assumptions"
assumptions["B2"] = "Subject Assumptions"
assumptions["B5"] = "Cap Rate"
assumptions["C5"] = 0.065
assumptions["C5"].number_format = PERCENT
assumptions["B6"] = "Subject Square Feet"
assumptions["C6"] = 36000
assumptions["C6"].number_format = "#,##0"
assumptions["B7"] = "Subject Units"
assumptions["C7"] = 24
assumptions["B8"] = "Market Vacancy"
assumptions["C8"] = 0.05
assumptions["C8"].number_format = PERCENT1

# ---- Pro Forma ---------------------------------------------------------
pro = wb.create_sheet("Pro Forma")
pro["B5"] = "Potential Gross Income"
pro["C5"] = 1500000
pro["C5"].number_format = CURRENCY
pro["B6"] = "Vacancy and Credit Loss"
pro["C6"] = "=-C5*Assumptions!C8"
pro["C6"].number_format = CURRENCY
pro["B7"] = "Effective Gross Income"
pro["C7"] = "=C5+C6"
pro["C7"].number_format = CURRENCY
pro["B8"] = "Operating Expenses"
pro["C8"] = 500000
pro["C8"].number_format = CURRENCY
pro["B9"] = "Net Operating Income"
pro["C9"] = "=C7-C8"
pro["C9"].number_format = CURRENCY
# A multi-column row: several figures on one line, so it must never be linked.
pro["B12"] = "Monthly Rent Roll"
pro["C12"] = 1000
pro["D12"] = 2000
pro["E12"] = 3000

# ---- Valuation ---------------------------------------------------------
valuation = wb.create_sheet("Valuation")
valuation["B5"] = "Indicated Value"
valuation["C5"] = "='Pro Forma'!C9/Assumptions!C5"
valuation["C5"].number_format = CURRENCY

# ---- Rent Roll (table) -------------------------------------------------
rent = wb.create_sheet("Rent Roll")
rent.append(["Unit", "Tenant", "Rent"])
rent.append(["101", "Alpha LLC", 1200])
rent.append(["102", "Beta Inc", 1350])
rent.append(["103", "Gamma Co", 1100])
table = Table(displayName="tblRentRoll", ref="A1:C4")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
rent.add_table(table)

# ---- Hidden sheet ------------------------------------------------------
calc = wb.create_sheet("Calc")
calc["B2"] = "Scratch Figure"
calc["C2"] = 12345
calc.sheet_state = "hidden"

# ---- Defined names -----------------------------------------------------
def add_name(name, ref, scope=None):
    dn = DefinedName(name, attr_text=ref)
    if scope is None:
        wb.defined_names[name] = dn
    else:
        scope.defined_names[name] = dn

add_name("Subject_Cap_Rate", "Assumptions!$C$5")
add_name("Subject_SF", "Assumptions!$C$6")
add_name("Subject_Units", "Assumptions!$C$7")
add_name("Subject_NOI", "'Pro Forma'!$C$9")
add_name("Subject_Value", "Valuation!$C$5")
add_name("Rent_Range", "'Rent Roll'!$C$2:$C$4")
add_name("Derived_Name", "Assumptions!$C$5*2")
# Same identifier defined again at sheet scope: must make Subject_NOI ambiguous
# for the workbook-level definition, so neither is trusted.
add_name("Subject_NOI", "Valuation!$C$5", scope=valuation)

wb.save(XLSX)

# Recalculate in place so every formula carries a cached result.
subprocess.run(["python3", RECALC, XLSX, "90"], check=True, cwd=os.path.dirname(RECALC))

# ---- Macro-enabled variant --------------------------------------------
# Same content, with a VBA project part and the macro-enabled content type.
# The VBA bytes are an inert placeholder: Keystone must never read or run them.
with zipfile.ZipFile(XLSX) as source, zipfile.ZipFile(XLSM, "w", zipfile.ZIP_DEFLATED) as target:
    for item in source.infolist():
        data = source.read(item.filename)
        if item.filename == "[Content_Types].xml":
            data = data.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            )
            data = data.replace(
                b"</Types>",
                b'<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
            )
        if item.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
            )
        target.writestr(item, data)
    target.writestr("xl/vbaProject.bin", b"INERT-PLACEHOLDER-NOT-EXECUTED")

print("built", XLSX, "and", XLSM)
